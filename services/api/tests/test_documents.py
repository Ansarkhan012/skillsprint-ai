from io import BytesIO
import asyncio
from pathlib import Path
from uuid import uuid4

import fitz
import pytest
from docx import Document
from fastapi import HTTPException
from fastapi.testclient import TestClient

from app import documents, security
from app.document_processing import (
    DOCX_MIME, PDF_MIME, DocumentInput, DocumentProblem, SourceUnit,
    chunk_units, parse_docx, parse_pdf, validate_file,
)
from app.document_repository import get_document_repository
from app.document_repository import DocumentRepository
from app.main import app
from app.config import get_settings
from app.models import AppRole, Profile
from app.supabase import get_gateway


def pdf_bytes(*pages: str) -> bytes:
    pdf = fitz.open()
    for text in pages:
        page = pdf.new_page()
        if text:
            page.insert_text((50, 50), text)
    result = pdf.tobytes()
    pdf.close()
    return result


def docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Safety", level=1)
    doc.add_heading("Daily checks", level=2)
    doc.add_paragraph("Wear protective equipment.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Step"
    table.cell(0, 1).text = "Check the equipment"
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def test_file_validation_pdf_docx_and_checksum():
    for filename, mime, data in (("manual.pdf", PDF_MIME, pdf_bytes("Hello")),
                                 ("manual.docx", DOCX_MIME, docx_bytes())):
        result_mime, checksum = validate_file(filename, mime, data, 15 * 1024 * 1024)
        assert result_mime == mime
        assert len(checksum) == 64
        assert checksum == validate_file(filename, mime, data, 15 * 1024 * 1024)[1]


@pytest.mark.parametrize("filename,mime,data,limit,code", [
    ("manual.txt", "text/plain", b"text", 100, "UNSUPPORTED_FILE_TYPE"),
    ("manual.pdf", PDF_MIME, b"", 100, "EMPTY_FILE"),
    ("manual.pdf", PDF_MIME, b"%PDF-abcdef", 5, "FILE_TOO_LARGE"),
    ("manual.pdf", DOCX_MIME, b"%PDF-abcdef", 100, "MIME_MISMATCH"),
    ("../manual.pdf", PDF_MIME, b"%PDF-abcdef", 100, "UNSAFE_FILENAME"),
    ("manual.docx", DOCX_MIME, b"PK-not-docx", 100, "MIME_MISMATCH"),
])
def test_file_validation_rejects_bad_inputs(filename, mime, data, limit, code):
    with pytest.raises(DocumentProblem) as exc:
        validate_file(filename, mime, data, limit)
    assert exc.value.code == code


def test_document_metadata_dates_and_codes():
    valid = DocumentInput(document_code="SAFETY_01", title="Safety", category="Policy",
                          version_label="v1", effective_date="2026-01-01", expiry_date="2026-12-31")
    assert valid.document_code == "SAFETY_01"
    with pytest.raises(ValueError):
        DocumentInput(document_code="bad code", title="Safety", category="Policy",
                      version_label="v1", effective_date="2026-01-01")
    with pytest.raises(ValueError):
        DocumentInput(document_code="SAFETY_01", title="Safety", category="Policy",
                      version_label="v1", effective_date="2026-12-31", expiry_date="2026-01-01")


def test_pdf_pages_and_scanned_fallback():
    units = parse_pdf(pdf_bytes("First page policy", "Second page instruction"))
    assert [unit.page_number for unit in units] == [1, 2]
    assert [unit.source_location["page"] for unit in units] == [1, 2]
    assert "First page" in units[0].text
    assert parse_pdf(pdf_bytes("")) == []


def test_docx_heading_paragraph_and_table_traceability():
    units = parse_docx(docx_bytes())
    paragraph = next(unit for unit in units if "protective" in unit.text)
    table_cell = next(unit for unit in units if "Check the equipment" in unit.text)
    assert paragraph.heading == "Daily checks"
    assert paragraph.section_path == "Safety / Daily checks"
    assert paragraph.source_location["paragraph"] > 0
    assert table_cell.source_location["table"] == 1
    assert table_cell.source_location["cell"] == 2


def test_chunks_stable_order_section_and_no_empty():
    version = uuid4()
    units = [SourceUnit("One two three four five six seven", {"kind": "pdf", "page": 2},
                        heading="Safety", section_path="Safety", page_number=2),
             SourceUnit("   ", {"kind": "pdf", "page": 3})]
    first = chunk_units(units, version, 14, 3)
    second = chunk_units(units, version, 14, 3)
    assert first == second
    assert [item["sequence"] for item in first] == list(range(len(first)))
    assert all(item["content"] and item["page_number"] == 2 for item in first)
    assert all(item["heading"] == "Safety" and item["source_location"]["page"] == 2 for item in first)


class FakeGateway:
    def __init__(self, profile):
        self.profile = profile

    async def get_profile(self, token, user_id):
        return self.profile


class FakeDocumentRepository:
    def __init__(self):
        self.calls = []
        self.storage_failure = False
        self.duplicate = False
        self.persistence_failure = False

    async def rpc(self, token, name, payload):
        self.calls.append((name, payload))
        if name == "begin_document_upload":
            if self.duplicate:
                raise HTTPException(status_code=409, detail="DOCUMENT_DUPLICATE")
            version_id = payload["p_version_id"]
            document_id = str(uuid4())
            return {"document_id": document_id, "version_id": version_id,
                    "storage_path": f"{document_id}/{version_id}.pdf"}
        if name == "finish_document_processing" and self.persistence_failure:
            raise HTTPException(status_code=503, detail="DOCUMENT_SERVICE_UNAVAILABLE")
        return None

    async def finalize_processing(self, payload):
        self.calls.append(("finish_document_processing", payload))
        if self.persistence_failure:
            raise HTTPException(status_code=503, detail="DOCUMENT_SERVICE_UNAVAILABLE")

    async def upload_original(self, token, path, mime_type, data):
        if self.storage_failure:
            raise HTTPException(status_code=503, detail="DOCUMENT_STORAGE_UNAVAILABLE")

    async def rows(self, token, table, params):
        if table == "document_versions" and params.get("select") == "storage_path,mime_type,original_filename":
            version_id = params["id"].removeprefix("eq.")
            return [{"storage_path": f"{uuid4()}/{version_id}.pdf", "mime_type": PDF_MIME,
                     "original_filename": "manual.pdf"}]
        return []

    async def download_original(self, token, path):
        return pdf_bytes("Original")


@pytest.fixture
def api_client(monkeypatch):
    repo = FakeDocumentRepository()
    test_profile = Profile(id=uuid4(), auth_user_id=uuid4(), display_name="Test",
                           status="ACTIVE", roles=frozenset({AppRole.TRAINING_MANAGER}))
    monkeypatch.setattr(security, "decode_access_token", lambda token, settings: test_profile.auth_user_id)
    app.dependency_overrides[get_gateway] = lambda: FakeGateway(test_profile)
    app.dependency_overrides[get_document_repository] = lambda: repo
    app.dependency_overrides[get_settings] = lambda: get_settings().model_copy(
        update={"supabase_service_role_key": "test-placeholder"})
    with TestClient(app) as client:
        yield client, repo, test_profile
    app.dependency_overrides.clear()


def upload(client, data=None):
    return client.post("/api/v1/documents/uploads", headers={"Authorization": "Bearer fake"},
        data={"document_code": "SAFETY_01", "title": "Safety", "category": "Policy",
              "version_label": "v1", "effective_date": "2026-01-01"},
        files={"file": ("manual.pdf", pdf_bytes("Safety first") if data is None else data, PDF_MIME)})


def test_upload_role_and_auth(api_client):
    client, repo, test_profile = api_client
    assert client.post("/api/v1/documents/uploads").status_code == 401
    test_profile.roles = frozenset({AppRole.EMPLOYEE})
    assert upload(client).status_code == 403
    test_profile.roles = frozenset({AppRole.TRAINING_MANAGER})
    response = upload(client)
    assert response.status_code == 202
    assert response.json()["parse_status"] == "PARSED"
    assert [name for name, _ in repo.calls] == ["begin_document_upload", "mark_document_processing", "finish_document_processing"]


def test_upload_invalid_and_duplicate(api_client):
    client, repo, _ = api_client
    assert upload(client, b"").json()["code"] == "EMPTY_FILE"
    assert not repo.calls
    repo.duplicate = True
    assert upload(client).status_code == 409


def test_duplicate_checksum_and_version_are_conflicts(api_client):
    client, repo, _ = api_client
    for conflict in ("checksum", "version_label"):
        repo.duplicate = True  # Simulates the migration's two unique constraints.
        response = upload(client)
        assert response.status_code == 409, conflict
        assert response.json()["code"] == "DOCUMENT_DUPLICATE"


def test_repository_maps_database_unique_conflict():
    import httpx

    response = httpx.Response(409, json={"code": "23505", "message": "duplicate"},
                              request=httpx.Request("POST", "https://example.test"))
    with pytest.raises(HTTPException) as exc:
        DocumentRepository.check(response)
    assert exc.value.status_code == 409


def test_storage_and_parser_failure_never_mark_parsed(api_client, monkeypatch):
    client, repo, _ = api_client
    repo.storage_failure = True
    assert upload(client).status_code == 503
    assert repo.calls[-1][1]["p_status"] == "FAILED"
    repo.calls.clear()
    repo.storage_failure = False
    monkeypatch.setattr(documents, "parse_pdf", lambda data: (_ for _ in ()).throw(DocumentProblem("PDF_PARSE_FAILED")))
    response = upload(client)
    assert response.status_code == 202
    assert response.json()["parse_status"] == "FAILED"
    assert repo.calls[-1][1]["p_status"] == "FAILED"


def test_persistence_failure_does_not_report_parsed(api_client):
    client, repo, _ = api_client
    repo.persistence_failure = True
    response = upload(client)
    assert response.status_code == 503
    assert repo.calls[-1][0] == "finish_document_processing"
    assert repo.calls[-1][1]["p_status"] == "PARSED"


def test_scanned_pdf_needs_review(api_client):
    client, repo, _ = api_client
    response = upload(client, pdf_bytes(""))
    assert response.status_code == 202
    assert response.json()["parse_status"] == "NEEDS_REVIEW"
    assert repo.calls[-1][1]["p_chunks"] == []


def test_review_role_separation(api_client):
    client, repo, test_profile = api_client
    path = f"/api/v1/document-versions/{uuid4()}"
    assert client.post(path + "/approve", headers={"Authorization": "Bearer fake"}, json={}).status_code == 403
    test_profile.roles = frozenset({AppRole.REVIEWER})
    assert client.get("/api/v1/documents", headers={"Authorization": "Bearer fake"}).status_code == 200
    assert client.post(path + "/submit", headers={"Authorization": "Bearer fake"}).status_code == 403
    assert client.post(path + "/reject", headers={"Authorization": "Bearer fake"}, json={}).status_code == 422
    assert client.post(path + "/approve", headers={"Authorization": "Bearer fake"}, json={}).status_code == 200
    assert repo.calls[-1][0] == "decide_document_version"


def test_document_reads_are_not_granted_to_employee(api_client):
    client, _, test_profile = api_client
    test_profile.roles = frozenset({AppRole.EMPLOYEE})
    headers = {"Authorization": "Bearer fake"}
    assert client.get("/api/v1/documents", headers=headers).status_code == 403
    assert client.get(f"/api/v1/documents/{uuid4()}", headers=headers).status_code == 403
    assert client.get(f"/api/v1/document-versions/{uuid4()}/chunks", headers=headers).status_code == 403
    assert client.get(f"/api/v1/document-versions/{uuid4()}/original", headers=headers).status_code == 403


def test_reviewer_can_download_private_original(api_client):
    client, _, test_profile = api_client
    test_profile.roles = frozenset({AppRole.REVIEWER})
    response = client.get(f"/api/v1/document-versions/{uuid4()}/original",
                          headers={"Authorization": "Bearer fake"})
    assert response.status_code == 200
    assert response.content.startswith(b"%PDF-")
    assert "attachment" in response.headers["content-disposition"]


def test_upload_invalid_metadata_returns_controlled_error(api_client):
    client, repo, _ = api_client
    response = client.post("/api/v1/documents/uploads", headers={"Authorization": "Bearer fake"},
        data={"document_code": "INVALID CODE", "title": "Safety", "category": "Policy",
              "version_label": "v1", "effective_date": "2026-12-31", "expiry_date": "2026-01-01"},
        files={"file": ("manual.pdf", pdf_bytes("Safety"), PDF_MIME)})
    assert response.status_code == 422
    assert response.json()["code"] == "INVALID_DOCUMENT_METADATA"
    assert not repo.calls


def test_heading_only_docx_preserves_source():
    doc = Document()
    doc.add_heading("Employees must complete safety training", level=1)
    stream = BytesIO()
    doc.save(stream)
    units = parse_docx(stream.getvalue())
    assert len(units) == 1
    assert units[0].text == "Employees must complete safety training"
    assert units[0].source_location["kind_detail"] == "heading"
    assert chunk_units(units, uuid4(), 1200, 150)[0]["paragraph_start"] == 1


def test_pdf_heading_and_source_block_position(monkeypatch):
    class Page:
        def get_text(self, *args, **kwargs):
            return {"blocks": [
                {"lines": [], "bbox": (0, 0, 0, 0)},
                {"lines": [{"spans": [{"text": "Safety policy requirement", "size": 20}]}], "bbox": (0, 10, 20, 20)},
                {"lines": [{"spans": [{"text": "All workers must comply with this policy every day", "size": 10}]}], "bbox": (0, 30, 20, 40)},
            ]}
    class Pdf:
        is_encrypted = False
        def __iter__(self):
            return iter([Page()])
        def close(self):
            pass
    monkeypatch.setattr("app.document_processing.fitz.open", lambda **kwargs: Pdf())
    units = parse_pdf(b"dummy")
    assert units[0].source_location["block"] == 2
    assert units[0].text == "Safety policy requirement"
    assert units[1].source_location["block"] == 3
    assert units[1].heading == units[0].text


def test_migration_blocks_user_finalization_and_dual_role_self_review():
    sql = (Path(__file__).parents[3] / "supabase/migrations/202609230002_document_intelligence.sql").read_text()
    finalizer = sql.split("create function public.finish_document_processing", 1)[1].split("end $$;", 1)[0]
    assert "auth.role()) is distinct from 'service_role'" in finalizer
    assert "grant execute on function public.finish_document_processing(uuid,public.document_parse_status,text,jsonb,text) to service_role" in sql
    assert "grant execute on function public.finish_document_processing(uuid,public.document_parse_status,text,jsonb,text) to authenticated" not in sql
    decision = sql.split("create function public.decide_document_version", 1)[1].split("end $$;", 1)[0]
    assert "v_version.uploaded_by = v_actor or v_version.submitted_by = v_actor" in decision
    assert "not public.has_app_role('ADMIN') or not p_admin_override" in decision
    assert "admin_override = v_override" in decision


def test_trusted_repository_uses_backend_credential_only_for_finalization(monkeypatch):
    import httpx
    captured = []
    class Client:
        def __init__(self, **kwargs):
            pass
        async def __aenter__(self):
            return self
        async def __aexit__(self, *args):
            pass
        async def post(self, url, headers, json):
            captured.append((url, headers, json))
            return httpx.Response(204, request=httpx.Request("POST", url))
    monkeypatch.setattr(httpx, "AsyncClient", Client)
    repo = DocumentRepository("https://example.test", "public-test", "backend-test")
    asyncio.run(repo.finalize_processing({"p_status": "PARSED"}))
    assert captured[0][0].endswith("/rpc/finish_document_processing")
    assert captured[0][1]["Authorization"] == "Bearer backend-test"
    assert "public-test" not in str(captured)


def test_phase2_sql_read_only_table_grants_and_rls():
    sql = (Path(__file__).parents[3] / "supabase/migrations/202609230002_document_intelligence.sql").read_text()
    for table in ("documents", "document_versions", "document_chunks"):
        assert f"alter table public.{table} enable row level security;" in sql
    assert "revoke all on public.documents, public.document_versions, public.document_chunks from public, anon, authenticated;" in sql
    assert "grant select on public.documents, public.document_versions, public.document_chunks to authenticated;" in sql
    assert "grant insert" not in sql and "grant update" not in sql and "grant delete" not in sql
    assert "security definer set search_path = ''" in sql
    assert "security invoker set search_path = ''" in sql


def test_documents_and_chunks_are_bounded_pages(api_client, monkeypatch):
    client, repo, _ = api_client
    async def rows(token, table, params):
        return [{"id": str(uuid4())} for _ in range(101)]
    monkeypatch.setattr(repo, "rows", rows)
    headers = {"Authorization": "Bearer fake"}
    for path in ("/api/v1/documents", f"/api/v1/document-versions/{uuid4()}/chunks"):
        response = client.get(path, headers=headers)
        assert response.status_code == 200
        assert len(response.json()["items"]) == 100
        assert response.json()["has_more"] is True
        assert client.get(path + "?limit=101", headers=headers).status_code == 422


def test_review_payload_exposes_explicit_override(api_client):
    client, repo, profile = api_client
    profile.roles = frozenset({AppRole.ADMIN})
    path = f"/api/v1/document-versions/{uuid4()}/approve"
    assert client.post(path, headers={"Authorization": "Bearer fake"},
                       json={"reason": "emergency", "admin_override": True}).status_code == 200
    assert repo.calls[-1][1]["p_admin_override"] is True


def test_oversized_multipart_rejected_before_route(api_client, monkeypatch):
    from types import SimpleNamespace
    from app import upload_limit
    client, repo, _ = api_client
    monkeypatch.setattr(upload_limit, "get_settings", lambda: SimpleNamespace(max_upload_bytes=1))
    response = client.post("/api/v1/documents/uploads", headers={"Authorization": "Bearer fake"},
                           files={"file": ("big.pdf", b"%PDF-" + b"x" * (1024 * 1024 + 1), PDF_MIME)})
    assert response.status_code == 413
    assert response.json()["code"] == "REQUEST_TOO_LARGE"
    assert not repo.calls


def test_request_limit_without_content_length(monkeypatch):
    from types import SimpleNamespace
    from app import upload_limit
    monkeypatch.setattr(upload_limit, "get_settings", lambda: SimpleNamespace(max_upload_bytes=1))
    reached_app = False
    sent = []
    chunks = iter([
        {"type": "http.request", "body": b"a" * 600000, "more_body": True},
        {"type": "http.request", "body": b"b" * 600000, "more_body": False},
    ])
    async def inner(scope, receive, send):
        nonlocal reached_app
        reached_app = True
    async def receive():
        return next(chunks)
    async def send(message):
        sent.append(message)
    scope = {"type": "http", "method": "POST", "path": "/api/v1/documents/uploads", "headers": []}
    asyncio.run(upload_limit.UploadRequestLimit(inner)(scope, receive, send))
    assert not reached_app
    assert sent[0]["status"] == 413


def test_retry_rehashes_existing_original(api_client, monkeypatch):
    from hashlib import sha256
    client, repo, profile = api_client
    version_id = uuid4()
    original = pdf_bytes("Retry source")
    async def rows(token, table, params):
        return [{"id": str(version_id), "document_id": str(uuid4()), "parse_status": "PROCESSING",
                 "review_status": "DRAFT", "uploaded_by": str(profile.id),
                 "storage_path": f"{uuid4()}/{version_id}.pdf", "sha256": sha256(original).hexdigest(),
                 "size_bytes": len(original), "mime_type": PDF_MIME, "original_filename": "retry.pdf"}]
    async def download(token, path):
        return original
    monkeypatch.setattr(repo, "rows", rows)
    monkeypatch.setattr(repo, "download_original", download)
    response = client.post(f"/api/v1/document-versions/{version_id}/retry",
                           headers={"Authorization": "Bearer fake"})
    assert response.status_code == 202
    assert response.json()["parse_status"] == "PARSED"
    assert [name for name, _ in repo.calls] == ["finish_document_processing"]


def test_malformed_pdf_never_parsed(api_client):
    client, repo, _ = api_client
    response = upload(client, b"%PDF-not-a-valid-document")
    assert response.status_code == 202
    assert response.json()["parse_status"] == "FAILED"
    assert repo.calls[-1][1]["p_chunks"] == []
